# ============================================================
# EXCEL HYBRID: TestSteps là template; TestData là meta + giá trị placeholder từng ca.
# Collection kiểm tra tất cả dữ liệu trước khi fixture mở browser.
# Không sort để che lỗi thứ tự step, không return [] để biến lỗi thành skip.
#
# TestSteps: template_id | step | keyword | args (mảng JSON) | note
# TestData : case_id | template_id | title | tags | requires | known_bug | <cột placeholder...>
#            tags / requires: nhiều giá trị ngăn bởi dấu phẩy, vd: "customer, allowWrite"
# ============================================================
from openpyxl import load_workbook

from utils.template_binding import (
    META_FIELDS,
    TEMPLATE_FIELDS,
    bind_template,
    parse_args,
    validate_template,
)

STEP_COLUMNS = ["template_id", "step", "keyword", "args", "note"]
DATA_COLUMNS = ["case_id", "template_id", "title", "tags", "requires", "known_bug"]


def _split(text):
    return [item.strip() for item in text.split(",") if item.strip()]


def build_cases(template_rows, data_rows):
    """Join theo template_id: một data row -> một case, không nhân tổ hợp."""
    if not template_rows or not data_rows:
        raise ValueError("TestSteps/TestData không được rỗng")
    templates = {}
    for row in template_rows:
        if not isinstance(row, dict) or set(row) != TEMPLATE_FIELDS:
            raise ValueError("TestSteps sai schema")
        templates.setdefault(row["template_id"], []).append(row)
    for template in templates.values():
        validate_template(template)

    cases, seen, referenced = [], set(), set()
    for row in data_rows:
        if not isinstance(row, dict) or not META_FIELDS <= set(row):
            raise ValueError("TestData sai schema")
        if any(not isinstance(value, str) for value in row.values()):
            raise ValueError("TestData cần giá trị chuỗi; không chấp nhận null")
        case_id, template_id = row["case_id"], row["template_id"]
        if not case_id or case_id in seen:
            raise ValueError(f"case_id rỗng hoặc trùng: {case_id!r}")
        if template_id not in templates:
            raise ValueError(f"Template không tồn tại: {template_id}")
        if not row["title"]:
            raise ValueError(f"Case {case_id}: thiếu title")
        seen.add(case_id)
        referenced.add(template_id)
        cases.append({
            "id": case_id,
            "template_id": template_id,
            "title": row["title"],
            "tags": _split(row["tags"]),
            "requires": _split(row["requires"]),
            "knownBug": row["known_bug"] or None,
            "steps": bind_template(templates[template_id], row),
        })
    unused = set(templates) - referenced
    if unused:
        raise ValueError(f"Template chưa có dữ liệu: {sorted(unused)}")
    return cases


def _cell(value):
    # Giữ input rỗng và số 0 đầu; không tự trim dữ liệu người dùng.
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_sheet(workbook, name, required, exact):
    if name not in workbook.sheetnames:
        raise ValueError(f"Thiếu sheet {name}")
    rows = list(workbook[name].iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Sheet {name} rỗng")
    header = [_cell(value) for value in rows[0]]
    if exact and header != required:
        raise ValueError(f"{name} sai cột: {header}; cần {required}")
    if not exact and header[: len(required)] != required:
        raise ValueError(f"{name} phải bắt đầu bằng các cột {required}; đang có {header}")
    if len(set(header)) != len(header) or "" in header:
        raise ValueError(f"{name} có cột trùng hoặc cột không tên")
    records = []
    for values in rows[1:]:
        if all(value is None for value in values):
            continue
        records.append({column: _cell(value) for column, value in zip(header, values)})
    return records


def read_hybrid_cases(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        steps = _read_sheet(workbook, "TestSteps", STEP_COLUMNS, exact=True)
        data = _read_sheet(workbook, "TestData", DATA_COLUMNS, exact=False)
    finally:
        workbook.close()
    for number, row in enumerate(steps, 2):
        row["args"] = parse_args(row["args"], f"{path.name} TestSteps dòng {number}")
    return build_cases(steps, data)
